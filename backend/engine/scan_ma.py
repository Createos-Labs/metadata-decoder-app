"""
M&A Audit Scanner
=================
Schema-free heuristic scanner for acquisition due-diligence files.
Unlike the LE ingestion scanner, this makes no assumptions about column names
or sheet structure. It scans every sheet in the workbook and flags data gaps
using heuristic rules.

Each finding has a stable ID (hash of sheet + check_type + field) so that
user-assigned severities survive re-scans when new files are added to the
acquisition.

Severity is NOT assigned here — that is done by the analyst in the app.
"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

import pandas as pd

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

ISRC_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{3}[0-9]{7}$")

PLACEHOLDER_PATTERNS = re.compile(
    r"\b(tbd|to be determined|placeholder|unknown|n/?a|pending|missing|none|null)\b",
    re.IGNORECASE,
)

# Column name fragments → "criticality hint" used when building Why-It-Matters text.
# Matching is case-insensitive substring.
CRITICAL_FIELD_HINTS: dict[str, str] = {
    "tax":          "Tax ID is legally required to issue payments in the US and most territories.",
    "ssn":          "SSN is legally required to issue 1099s to US-based payees.",
    "ein":          "EIN is legally required to issue 1099s to US-based entities.",
    "isrc":         "ISRC is the unique identifier for each recording. Without it, streaming revenue cannot be accurately attributed or migrated.",
    "iswc":         "ISWC is the unique identifier for each musical work. Required for publishing royalty reporting.",
    "upc":          "UPC/EAN is the unique product identifier. Required for DSP catalog delivery and revenue matching.",
    "email":        "Email is required for payment portal onboarding and statement delivery.",
    "term-start":   "Contract start date is needed to verify whether deals are active and enforceable.",
    "term-end":     "Contract end date is needed to verify whether deals are active and enforceable.",
    "start-date":   "Contract start date is needed to verify whether deals are active and enforceable.",
    "end-date":     "Contract end date is needed to verify whether deals are active and enforceable.",
    "release-date": "Release date is required for catalog reporting and DSP metadata delivery.",
    "c-line":       "C-line (copyright) is required metadata for DSP ingestion.",
    "p-line":       "P-line (phonogram producer) is required metadata for DSP ingestion.",
    "genre":        "Genre is required for many DSP catalog submissions.",
    "address":      "Physical address is required for payees receiving manual or check payments and for tax documentation.",
    "payee":        "Payee information is required to route royalty payments correctly.",
    "label":        "Label name must be consistent — inconsistencies cause ingestion failures.",
    "artist":       "Artist name is required for catalog attribution and royalty allocation.",
    "title":        "Track/work title is required for catalog identification.",
    "composer":     "Composer credits are needed for mechanical royalty reporting.",
    "writer":       "Writer credits are needed for publishing royalty reporting.",
    "split":        "Split percentages determine how royalties are divided. Missing or incorrect splits block payment.",
    "rate":         "Royalty rate is required to calculate payments.",
    "price":        "Price fields are required to calculate percentage-of-wholesale royalties.",
}

# Columns whose absence is especially critical (used for higher-priority warnings).
BLOCKER_FIELD_FRAGMENTS = {"tax", "ssn", "ein", "isrc"}
HIGH_FIELD_FRAGMENTS = {"email", "term-start", "term-end", "start-date", "end-date",
                        "c-line", "p-line", "upc", "iswc", "payee", "artist", "rate"}

# % missing thresholds for generic (non-named-critical) columns.
COMPLETENESS_HIGH_THRESHOLD = 0.90   # >90% blank on an otherwise important column
COMPLETENESS_MEDIUM_THRESHOLD = 0.50  # >50% blank

# Minimum rows before we bother with completeness checks (skip near-empty sheets).
MIN_ROWS_FOR_CHECKS = 3


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _finding_id(sheet: str, check_type: str, field: str) -> str:
    """Stable ID so severity assignments survive re-scans."""
    key = f"{sheet}|{check_type}|{field}"
    return hashlib.sha1(key.encode()).hexdigest()[:12]


def _col_hint(col: str) -> str | None:
    """Return the hint text if the column name matches a critical field fragment."""
    col_lower = col.lower().replace(" ", "-").replace("_", "-")
    for fragment, hint in CRITICAL_FIELD_HINTS.items():
        if fragment in col_lower:
            return hint
    return None


def _is_isrc_column(col: str) -> bool:
    col_lower = col.lower().replace(" ", "").replace("-", "").replace("_", "")
    return "isrc" in col_lower


def _is_blank(val) -> bool:
    if val is None:
        return True
    if isinstance(val, float):
        import math
        return math.isnan(val)
    return str(val).strip() == ""


def _stringify(val) -> str:
    if _is_blank(val):
        return ""
    return str(val).strip()


# ---------------------------------------------------------------------------
# Per-sheet checks
# ---------------------------------------------------------------------------

def _check_completeness(
    df: pd.DataFrame,
    sheet: str,
    findings: list[dict],
) -> None:
    """Flag columns with significant missing data."""
    n = len(df)
    if n < MIN_ROWS_FOR_CHECKS:
        return

    for col in df.columns:
        blank_count = sum(1 for v in df[col] if _is_blank(v))
        pct_missing = blank_count / n

        if pct_missing == 0:
            continue

        col_lower = col.lower().replace(" ", "-").replace("_", "-")
        hint = _col_hint(col)
        is_critical = any(f in col_lower for f in BLOCKER_FIELD_FRAGMENTS)
        is_high = any(f in col_lower for f in HIGH_FIELD_FRAGMENTS)

        # Decide whether to surface this gap.
        if blank_count == n:
            # 100% missing — always surface.
            severity_hint = "BLOCKER" if is_critical else ("HIGH" if is_high else "MEDIUM")
            label = "100% of rows"
        elif is_critical and pct_missing > 0.05:
            severity_hint = "BLOCKER"
            label = f"{blank_count:,} of {n:,} rows ({pct_missing:.0%})"
        elif is_high and pct_missing >= COMPLETENESS_MEDIUM_THRESHOLD:
            severity_hint = "HIGH"
            label = f"{blank_count:,} of {n:,} rows ({pct_missing:.0%})"
        elif pct_missing >= COMPLETENESS_HIGH_THRESHOLD:
            severity_hint = "MEDIUM"
            label = f"{blank_count:,} of {n:,} rows ({pct_missing:.0%})"
        elif pct_missing >= COMPLETENESS_MEDIUM_THRESHOLD and hint:
            severity_hint = "MEDIUM"
            label = f"{blank_count:,} of {n:,} rows ({pct_missing:.0%})"
        else:
            continue

        why = hint or f"This field is {pct_missing:.0%} empty across {n:,} rows."

        findings.append({
            "id": _finding_id(sheet, "completeness", col),
            "sheet": sheet,
            "field": col,
            "check_type": "completeness",
            "title": f"{col} — missing on {label}",
            "finding": f"Missing on {label}",
            "why_it_matters": why,
            "detail": {
                "missing_count": blank_count,
                "total_rows": n,
                "pct_missing": round(pct_missing * 100, 1),
            },
            "severity": None,
            "dismissed": False,
            "_severity_hint": severity_hint,
        })


def _check_isrc(
    df: pd.DataFrame,
    sheet: str,
    findings: list[dict],
) -> dict[str, str]:
    """
    Validate ISRC columns. Returns {isrc_value: sheet} map for cross-sheet checks.
    """
    isrc_map: dict[str, str] = {}
    n = len(df)
    if n < MIN_ROWS_FOR_CHECKS:
        return isrc_map

    for col in df.columns:
        if not _is_isrc_column(col):
            continue

        bad_format: list[str] = []
        valid_isrcs: list[str] = []

        for v in df[col]:
            if _is_blank(v):
                continue
            raw = _stringify(v).upper().replace("-", "").replace(" ", "")
            if ISRC_RE.match(raw):
                valid_isrcs.append(raw)
                isrc_map[raw] = sheet
            else:
                bad_format.append(_stringify(v))

        if bad_format:
            findings.append({
                "id": _finding_id(sheet, "isrc_format", col),
                "sheet": sheet,
                "field": col,
                "check_type": "isrc_format",
                "title": f"{col} — {len(bad_format)} malformed ISRC{'s' if len(bad_format) != 1 else ''}",
                "finding": f"{len(bad_format)} values don't match ISRC format (CCRRRYYNNNNN, 12 chars)",
                "why_it_matters": "Malformed ISRCs will be rejected during ingestion and break revenue attribution.",
                "detail": {"bad_count": len(bad_format), "examples": bad_format[:5]},
                "severity": None,
                "dismissed": False,
                "_severity_hint": "HIGH",
            })

        # Duplicate ISRC check (same ISRC on rows with different titles).
        title_col = next(
            (c for c in df.columns if "title" in c.lower() and "track" in c.lower()),
            next((c for c in df.columns if "title" in c.lower()), None),
        )
        if title_col and valid_isrcs:
            isrc_to_titles: dict[str, set[str]] = {}
            for _, row in df.iterrows():
                raw_isrc = _stringify(row.get(col, "")).upper().replace("-", "").replace(" ", "")
                if not ISRC_RE.match(raw_isrc):
                    continue
                title = _stringify(row.get(title_col, ""))
                if raw_isrc not in isrc_to_titles:
                    isrc_to_titles[raw_isrc] = set()
                if title:
                    isrc_to_titles[raw_isrc].add(title)

            conflicts = {k: v for k, v in isrc_to_titles.items() if len(v) > 1}
            if conflicts:
                examples = [
                    f"{isrc}: {' / '.join(sorted(titles)[:3])}"
                    for isrc, titles in list(conflicts.items())[:3]
                ]
                findings.append({
                    "id": _finding_id(sheet, "isrc_duplicate", col),
                    "sheet": sheet,
                    "field": col,
                    "check_type": "isrc_duplicate",
                    "title": f"{col} — {len(conflicts)} ISRC{'s' if len(conflicts) != 1 else ''} linked to multiple titles",
                    "finding": f"{len(conflicts)} ISRCs appear on rows with different track titles",
                    "why_it_matters": "The same ISRC should always correspond to the same recording. Conflicts suggest data errors that will cause mis-attribution.",
                    "detail": {"conflict_count": len(conflicts), "examples": examples},
                    "severity": None,
                    "dismissed": False,
                    "_severity_hint": "HIGH",
                })

    return isrc_map


def _check_placeholders(
    df: pd.DataFrame,
    sheet: str,
    findings: list[dict],
) -> None:
    """Flag name/payee columns containing placeholder values."""
    n = len(df)
    if n < MIN_ROWS_FOR_CHECKS:
        return

    name_cols = [
        c for c in df.columns
        if any(frag in c.lower() for frag in ("name", "payee", "artist", "client", "writer"))
    ]

    for col in name_cols:
        hits = []
        for v in df[col]:
            if _is_blank(v):
                continue
            s = _stringify(v)
            if PLACEHOLDER_PATTERNS.search(s):
                hits.append(s)

        if not hits:
            continue

        unique_hits = list(dict.fromkeys(hits))[:5]
        findings.append({
            "id": _finding_id(sheet, "placeholder", col),
            "sheet": sheet,
            "field": col,
            "check_type": "placeholder",
            "title": f"{col} — {len(hits)} placeholder/TBD {'values' if len(hits) != 1 else 'value'}",
            "finding": f"{len(hits)} rows have placeholder values (e.g. {', '.join(repr(h) for h in unique_hits[:3])})",
            "why_it_matters": "Placeholder payees and artist names cannot receive payments. These deals may be unfinalized.",
            "detail": {"hit_count": len(hits), "examples": unique_hits},
            "severity": None,
            "dismissed": False,
            "_severity_hint": "HIGH",
        })


def _check_status_inconsistencies(
    df: pd.DataFrame,
    sheet: str,
    findings: list[dict],
) -> None:
    """Flag INACTIVE payors/statuses that appear alongside active records."""
    n = len(df)
    if n < MIN_ROWS_FOR_CHECKS:
        return

    status_cols = [c for c in df.columns if "status" in c.lower() or "payor" in c.lower() or "payer" in c.lower()]

    for col in status_cols:
        inactive_count = sum(
            1 for v in df[col]
            if not _is_blank(v) and "inactive" in _stringify(v).lower()
        )
        active_count = sum(
            1 for v in df[col]
            if not _is_blank(v) and "active" in _stringify(v).lower()
            and "inactive" not in _stringify(v).lower()
        )

        if inactive_count > 0 and active_count > 0:
            findings.append({
                "id": _finding_id(sheet, "status_mix", col),
                "sheet": sheet,
                "field": col,
                "check_type": "status_mix",
                "title": f"{col} — {inactive_count} inactive alongside {active_count} active records",
                "finding": f"{inactive_count} INACTIVE records mixed with {active_count} ACTIVE records",
                "why_it_matters": "Inactive contracts or payors may still have catalog earning under their name. Confirm whether these are truly dead or should be migrated.",
                "detail": {"inactive_count": inactive_count, "active_count": active_count},
                "severity": None,
                "dismissed": False,
                "_severity_hint": "MEDIUM",
            })

        # Standalone INACTIVE payor (no active counterpart — could still be a flag).
        if inactive_count > 0 and active_count == 0:
            findings.append({
                "id": _finding_id(sheet, "all_inactive", col),
                "sheet": sheet,
                "field": col,
                "check_type": "all_inactive",
                "title": f"{col} — {inactive_count} INACTIVE {'rows' if inactive_count != 1 else 'row'}",
                "finding": f"{inactive_count} rows have status INACTIVE",
                "why_it_matters": "Inactive records may still have catalog earning. Confirm whether these should be excluded from migration.",
                "detail": {"inactive_count": inactive_count},
                "severity": None,
                "dismissed": False,
                "_severity_hint": "LOW",
            })


def _check_label_consistency(
    df: pd.DataFrame,
    sheet: str,
    findings: list[dict],
) -> None:
    """Flag columns that look like label names but have too many distinct values."""
    n = len(df)
    if n < MIN_ROWS_FOR_CHECKS:
        return

    label_cols = [c for c in df.columns if "label" in c.lower()]
    for col in label_cols:
        values = [_stringify(v) for v in df[col] if not _is_blank(v)]
        unique_vals = list(dict.fromkeys(values))
        if len(unique_vals) > 2:
            findings.append({
                "id": _finding_id(sheet, "label_inconsistency", col),
                "sheet": sheet,
                "field": col,
                "check_type": "label_inconsistency",
                "title": f"{col} — {len(unique_vals)} distinct values",
                "finding": f"{len(unique_vals)} different label name variants: {', '.join(unique_vals[:5])}",
                "why_it_matters": "LE requires a consistent legal entity name per asset. Inconsistent label names will cause ingestion failures and split attribution errors.",
                "detail": {"unique_count": len(unique_vals), "values": unique_vals[:10]},
                "severity": None,
                "dismissed": False,
                "_severity_hint": "HIGH",
            })


# ---------------------------------------------------------------------------
# Cross-sheet checks
# ---------------------------------------------------------------------------

def _check_cross_sheet_isrcs(
    isrc_maps: dict[str, dict[str, str]],
    findings: list[dict],
) -> None:
    """
    For sheets that both have ISRC columns, flag ISRCs that appear in one
    but not the other.
    """
    sheet_names = [s for s, m in isrc_maps.items() if m]
    if len(sheet_names) < 2:
        return

    all_isrcs: set[str] = set()
    for m in isrc_maps.values():
        all_isrcs.update(m.keys())

    for sheet_a in sheet_names:
        for sheet_b in sheet_names:
            if sheet_a >= sheet_b:
                continue
            only_in_a = set(isrc_maps[sheet_a]) - set(isrc_maps[sheet_b])
            only_in_b = set(isrc_maps[sheet_b]) - set(isrc_maps[sheet_a])

            if only_in_a:
                findings.append({
                    "id": _finding_id(f"{sheet_a}+{sheet_b}", "isrc_orphan", sheet_a),
                    "sheet": sheet_a,
                    "field": "ISRC",
                    "check_type": "isrc_orphan",
                    "title": f"{len(only_in_a)} ISRC{'s' if len(only_in_a) != 1 else ''} in '{sheet_a}' not found in '{sheet_b}'",
                    "finding": f"{len(only_in_a)} ISRCs exist in '{sheet_a}' but have no match in '{sheet_b}'",
                    "why_it_matters": "ISRCs present in one sheet but absent from another suggest incomplete data or transcription errors that will break cross-sheet reconciliation.",
                    "detail": {
                        "count": len(only_in_a),
                        "examples": sorted(only_in_a)[:5],
                        "sheet_a": sheet_a,
                        "sheet_b": sheet_b,
                    },
                    "severity": None,
                    "dismissed": False,
                    "_severity_hint": "MEDIUM",
                })

            if only_in_b:
                findings.append({
                    "id": _finding_id(f"{sheet_a}+{sheet_b}", "isrc_orphan", sheet_b),
                    "sheet": sheet_b,
                    "field": "ISRC",
                    "check_type": "isrc_orphan",
                    "title": f"{len(only_in_b)} ISRC{'s' if len(only_in_b) != 1 else ''} in '{sheet_b}' not found in '{sheet_a}'",
                    "finding": f"{len(only_in_b)} ISRCs exist in '{sheet_b}' but have no match in '{sheet_a}'",
                    "why_it_matters": "ISRCs present in one sheet but absent from another suggest incomplete data or transcription errors that will break cross-sheet reconciliation.",
                    "detail": {
                        "count": len(only_in_b),
                        "examples": sorted(only_in_b)[:5],
                        "sheet_a": sheet_b,
                        "sheet_b": sheet_a,
                    },
                    "severity": None,
                    "dismissed": False,
                    "_severity_hint": "MEDIUM",
                })


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def _load_sheets(input_path: Path) -> list[tuple[str, pd.DataFrame]]:
    """
    Load all data sheets from an Excel or CSV file.
    Returns a list of (sheet_name, DataFrame) tuples.
    CSV files are treated as a single sheet named after the file stem.
    """
    suffix = input_path.suffix.lower()

    if suffix == ".csv":
        # Try common encodings so mis-encoded exports don't crash.
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                df = pd.read_csv(input_path, header=0, encoding=enc, low_memory=False)
                df = df.dropna(how="all", axis=1).dropna(how="all", axis=0)
                df.columns = [str(c).strip() for c in df.columns]
                return [(input_path.stem, df)]
            except UnicodeDecodeError:
                continue
        return []

    # Excel: load sheet names without reading data first.
    from openpyxl import load_workbook as _lw
    wb = _lw(input_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    SKIP_SHEET_HINTS = {"dropdown", "dropdowns", "notes", "instructions", "legend", "ref", "lookup"}
    sheets = []
    for sheet_name in sheet_names:
        if any(hint in sheet_name.lower() for hint in SKIP_SHEET_HINTS):
            continue
        try:
            df = pd.read_excel(input_path, sheet_name=sheet_name, header=0)
            df = df.dropna(how="all", axis=1).dropna(how="all", axis=0)
            df.columns = [str(c).strip() for c in df.columns]
            sheets.append((sheet_name, df))
        except Exception:
            continue
    return sheets


def analyze(input_path: Path) -> dict:
    """
    Scan all sheets in the workbook (or a single CSV) and return a findings dict.

    Returns:
        {
          "sheets_scanned": [...],
          "findings": [...],          # full list, severity=None until analyst assigns
          "stats": {
              "total_findings": N,
              "sheets": {"SheetName": {"rows": N, "findings": N}, ...}
          }
        }
    """
    sheets = _load_sheets(input_path)

    findings: list[dict] = []
    isrc_maps: dict[str, dict[str, str]] = {}
    sheet_stats: dict[str, dict] = {}

    for sheet_name, df in sheets:

        # Drop fully-empty columns and rows.
        df = df.dropna(how="all", axis=1).dropna(how="all", axis=0)
        df.columns = [str(c).strip() for c in df.columns]

        n_rows = len(df)
        before = len(findings)

        if n_rows >= MIN_ROWS_FOR_CHECKS:
            _check_completeness(df, sheet_name, findings)
            isrc_map = _check_isrc(df, sheet_name, findings)
            isrc_maps[sheet_name] = isrc_map
            _check_placeholders(df, sheet_name, findings)
            _check_status_inconsistencies(df, sheet_name, findings)
            _check_label_consistency(df, sheet_name, findings)

        sheet_stats[sheet_name] = {
            "rows": n_rows,
            "columns": list(df.columns),
            "findings": len(findings) - before,
        }

    # Cross-sheet ISRC check.
    _check_cross_sheet_isrcs(isrc_maps, findings)

    return {
        "sheets_scanned": list(sheet_stats.keys()),
        "findings": findings,
        "stats": {
            "total_findings": len(findings),
            "sheets": sheet_stats,
        },
    }
